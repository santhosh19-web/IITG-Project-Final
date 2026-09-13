import os
import glob
try:
    from pptx import Presentation
    from pptx.util import Inches, Pt
except ImportError:
    pass

header_text = """# Name: K.Santhosh Reddy
# Cohort: 4 Batch A2
# Mentor: Mr. Vamshi

"""

def should_update_file(filename):
    if "Day_" in filename:
        try:
            parts = filename.split("Day_")[1].split("_")
            day_str = parts[0].split(".")[0]
            day_num = int(day_str)
            return 4 <= day_num <= 21
        except:
            pass
    
    # Specific edge case files for Day >= 4
    if "generate_review2_ppt.py" in filename: return True # Day 9
    if "WEEK_3_ML_Assisted_Decoder.py" in filename: return True # Day 18
    
    return False

def main():
    print("Updating python scripts...")
    py_files = glob.glob("*.py")
    for py_file in py_files:
        if py_file == "update_names.py":
            continue
        if should_update_file(py_file):
            with open(py_file, 'r', encoding='utf-8') as f:
                content = f.read()
            if "# Name: K.Santhosh Reddy" not in content:
                with open(py_file, 'w', encoding='utf-8') as f:
                    f.write(header_text + content)
                print(f"  [x] Updated {py_file}")

    print("\nUpdating PowerPoint files...")
    pptx_files = glob.glob("*.pptx")
    for pptx_file in pptx_files:
        if should_update_file(pptx_file):
            try:
                prs = Presentation(pptx_file)
                slide = prs.slides[0]
                
                left = Inches(0.5)
                top = Inches(0.2)
                width = Inches(5.0)
                height = Inches(1.0)
                txBox = slide.shapes.add_textbox(left, top, width, height)
                tf = txBox.text_frame
                tf.text = "Name: K.Santhosh Reddy\nCohort: 4 Batch A2\nMentor: Mr. Vamshi"
                
                prs.save(pptx_file)
                print(f"  [x] Updated {pptx_file}")
            except Exception as e:
                print(f"  [!] Failed to update {pptx_file}: {e}")

if __name__ == "__main__":
    main()
