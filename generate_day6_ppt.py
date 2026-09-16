from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

def set_font(run, font_name="Times New Roman", font_size=18, bold=False):
    run.font.name = font_name
    run.font.size = Pt(font_size)
    run.font.bold = bold

def create_slide(prs, title, bullets):
    slide_layout = prs.slide_layouts[1] # Title and Content
    slide = prs.slides.add_slide(slide_layout)
    
    # Title
    title_shape = slide.shapes.title
    title_shape.text = title
    set_font(title_shape.text_frame.paragraphs[0].runs[0], font_size=32, bold=True)
    
    # Body
    body_shape = slide.placeholders[1]
    tf = body_shape.text_frame
    tf.clear() # clear default paragraphs
    
    for text in bullets:
        p = tf.add_paragraph()
        p.text = text
        p.level = 0
        p.space_after = Pt(14)
        for run in p.runs:
            set_font(run, font_size=20)
            
    return slide

def main():
    prs = Presentation()
    
    # Slide 1: Title
    slide_layout = prs.slide_layouts[0]
    slide = prs.slides.add_slide(slide_layout)
    title = slide.shapes.title
    subtitle = slide.placeholders[1]
    
    title.text = "Day 6: Methodology V1 & Data-Readiness"
    set_font(title.text_frame.paragraphs[0].runs[0], font_size=40, bold=True)
    
    subtitle.text = "Name: K.Santhosh Reddy\nCohort: 4 Batch A2\nMentor: Mr. Vamshi"
    for p in subtitle.text_frame.paragraphs:
        for run in p.runs:
            set_font(run, font_size=24)
            
    # Slide 2: Literature & Data-Readiness
    bullets_2 = [
        "Literature Foundation: Builds on robust CV Region of Interest (ROI) extraction and Signal Processing.",
        "Constraint Adherence: Offline software pipelines provide the highest reliability for OOK modulation without physical hardware integration.",
        "Data Readiness: Raw dataset includes video sequences and a ground-truth mapping file.",
        "Validation: Videos strictly checked for Nyquist compliance (FPS > 2x baud rate).",
        "Methodology: Strict Group-by-Video split enforced to prevent temporal data leakage."
    ]
    create_slide(prs, "Literature & Data-Readiness", bullets_2)
    
    # Slide 3: Preprocessing & Modelling
    bullets_3 = [
        "Preprocessing (CV) - Frame Extraction: Read sequence frame-by-frame via OpenCV.",
        "Preprocessing (CV) - Localization: Identify the brightest stable region (LED) and extract a bounding box ROI.",
        "Preprocessing (CV) - Intensity Extraction: Calculate mean pixel intensity across ROI to generate a 1D time-series.",
        "Modelling (Baseline): Apply dynamic thresholding (Moving Average) to the 1D signal to classify ON and OFF states.",
        "Modelling (Advanced ML): Extract local statistical features and train SVM for noise resilience."
    ]
    create_slide(prs, "Proposed Preprocessing & Modelling Plan", bullets_3)
    
    # Slide 4: Evaluation & Constraints
    bullets_4 = [
        "Timing Recovery: Synchronize the classified symbols to account for camera frame-drops.",
        "Reconstruction: Group binary sequences and map them to standard character encodings (ASCII).",
        "Primary Metrics: Bit Error Rate (BER), Exact Decoded-Message Correctness, and execution time.",
        "Secondary Metric: Synchronization Robustness (maximum allowable frame drift).",
        "Expected Constraints: Models must mitigate fluctuating ambient light and rolling-shutter artifacts offline."
    ]
    create_slide(prs, "Evaluation Strategy & Constraints", bullets_4)
    
    # Slide 5: Split Strategy & Mentor Updates
    bullets_5 = [
        "Split Strategy: 70% Train, 15% Validation, 15% Test. Split is performed entirely at the video level.",
        "Mentor Updates - Parameterization: Moving average window size is dynamically parameterized for varying camera FPS.",
        "Mentor Updates - Metrics: Explicitly integrated maximum frame drift metric into the evaluation pipeline.",
        "Status Check: Methodology V1 is finalized and the theoretical design phase is officially frozen."
    ]
    create_slide(prs, "Split Strategy & Mentor Corrections", bullets_5)
    
    output_path = "WEEK_1_Day_6_Data_Capture.pptx"
    prs.save(output_path)
    print(f"Successfully generated {output_path}")

if __name__ == "__main__":
    main()
