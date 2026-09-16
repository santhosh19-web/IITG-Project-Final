import glob
from pptx import Presentation

def should_update_file(filename):
    if "Day_" in filename:
        try:
            parts = filename.split("Day_")[1].split("_")
            day_str = parts[0].split(".")[0]
            day_num = int(day_str)
            return day_num >= 4
        except:
            pass
    return False

def move_slide(prs, old_index, new_index):
    xml_slides = prs.slides._sldIdLst
    slides = list(xml_slides)
    xml_slides.remove(slides[old_index])
    xml_slides.insert(new_index, slides[old_index])

def main():
    pptx_files = glob.glob("*.pptx")
    for file in pptx_files:
        if should_update_file(file):
            try:
                prs = Presentation(file)
                
                # Find the index of the "Project Details" slide
                details_idx = -1
                for i, slide in enumerate(prs.slides):
                    if slide.shapes.title and slide.shapes.title.text == "Project Details":
                        details_idx = i
                        
                if details_idx != -1 and details_idx != 0:
                    # Move to index 0 (very first slide)
                    move_slide(prs, details_idx, 0)
                    prs.save(file)
                    print(f"Moved details slide to very front in {file}")
                else:
                    print(f"No changes needed for {file}")
                    
            except Exception as e:
                print(f"Error processing {file}: {e}")

if __name__ == "__main__":
    main()
