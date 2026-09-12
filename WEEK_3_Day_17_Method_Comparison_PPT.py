import os

try:
    import collections.abc
    from pptx import Presentation
    from pptx.util import Inches, Pt
except ImportError:
    print("Error: python-pptx is not installed. Please run 'pip install python-pptx'")
    exit(1)

def create_ppt():
    print("--- Week 3 Day 17: Method Comparison & Progress Pack ---")
    prs = Presentation()
    
    # Slide 1: Method Comparison Table
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    title = slide.shapes.title
    title.text = "Method Comparison (Static vs Adaptive vs Dynamic Phase)"
    
    # Create Table
    rows, cols = 4, 5
    left = Inches(0.5)
    top = Inches(1.5)
    width = Inches(9.0)
    height = Inches(2.0)
    table = slide.shapes.add_table(rows, cols, left, top, width, height).table
    
    headers = ["Method", "Decoded Correctness", "Robustness", "Execution Time", "Verdict"]
    for i, h in enumerate(headers):
        table.cell(0, i).text = h
        
    data = [
        ["Static Threshold (Day 11)", "Low (fails on drift)", "Poor (drifts)", "Fastest", "Rejected"],
        ["Adaptive Threshold (Day 15)", "Medium (fixes baseline)", "Medium (clock drifts)", "Fast", "Baseline"],
        ["Dynamic Sync (Day 16)", "High", "Robust (zero-cross)", "Medium", "Shortlisted"]
    ]
    
    for row_idx, row_data in enumerate(data):
        for col_idx, text in enumerate(row_data):
            table.cell(row_idx + 1, col_idx).text = text
            
    # Slide 2: Discussion
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    title = slide.shapes.title
    title.text = "Discussion: Overfitting, Limitations & Failures"
    content = slide.placeholders[1]
    content.text = (
        "• Overfitting: Tuning the threshold window size (50 frames) heavily to the 1LED video might overfit to its specific lighting drift.\n"
        "• Data Limitations: Lack of an explicit ground_truth.csv means BER is estimated via internal consistency rather than exact bit matches.\n"
        "• Complexity: Dynamic phase snapping adds O(N) lookahead searches per frame, increasing processing time compared to static sampling.\n"
        "• Failure Patterns: Blurring from rapid camera motion causes missed zero-crossings, forcing the decoder to dead-reckon."
    )
    
    # Slide 3: Shortlisted Candidates
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    title = slide.shapes.title
    title.text = "Shortlisted Candidates for Deeper Experiments"
    content = slide.placeholders[1]
    content.text = (
        "Candidate 1: Dynamic Sync Decoder (Adaptive Threshold + Zero-Cross Snapping)\n"
        " - Rationale: Only method currently capable of extracting 10,000+ stable bits from the drifting 1LED video.\n\n"
        "Future Enhancements for Candidate:\n"
        " - Implement a Phase-Locked Loop (PLL) to handle missed zero-crossings gracefully instead of dead-reckoning.\n"
        " - Test multi-channel (RGB) thresholding rather than pure grayscale intensity to improve resilience against color temperature changes."
    )
    
    output_filename = "WEEK_3_Day_17_Method_Comparison.pptx"
    prs.save(output_filename)
    print(f"Successfully generated 3-slide progress pack: {output_filename}")

if __name__ == "__main__":
    create_ppt()
